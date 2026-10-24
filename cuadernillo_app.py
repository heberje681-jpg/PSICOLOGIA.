import io
import re
import time
import unicodedata
import zlib
from urllib.parse import quote

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # falta python-docx en requirements.txt
    Document = None

st.set_page_config(page_title="Cuadernillo de exploración educativa", page_icon="📋", layout="centered")
try:
    PIN = str(st.secrets["PIN"])
except Exception:
    PIN = "1234"  # cámbialo aquí o en Secrets de Streamlit: PIN = "tu_pin"

# Umbrales HEURÍSTICOS (no validados). Se pueden editar en la pestaña Resultados.
U = {"omisiones": 8, "falsos_pos": 3, "recuerdo": 2, "instrucciones": 3, "comprension": 2, "mates": 4, "org_items": 4, "fatiga": 3}
ULAB = {"omisiones": "Omisiones act. 1 (≥)", "falsos_pos": "Falsos positivos act. 1 (≥)",
        "recuerdo": "Palabras recordadas (≤)", "instrucciones": "Instrucciones 2B (≤)",
        "comprension": "Comprensión lectora (≤)", "mates": "Matemáticas (≤)",
        "org_items": "Conductas de organización frecuentes (≥)",
        "fatiga": "Caída de aciertos 1ª vs 2ª mitad (≥)"}

FILAS = [f.split() for f in [
    "B A D C A F E A B C D A", "A C B D E A C F A D B E", "D E A B C D A E F B A C", "C A E D B A F C D A E B",
    "B D C A E F A B D C A E", "A E D B A C E D A F B C", "D A F C B A E D C A B E", "E C A D F B A C E D A B"]]
PALABRAS = ["luna", "zapato", "jardín", "moneda", "ventana"]
ACCIONES = ["Nombre arriba", "Círculo abajo a la derecha", "Número 8 dentro del círculo",
            "Nombre subrayado", "ESCUELA debajo del círculo"]
TEXTO = ("**EL PROYECTO ESCOLAR.** Un grupo de estudiantes decidió crear un pequeño huerto en la escuela. "
         "Antes de comenzar, investigaron qué plantas podían crecer en su región y cuánta agua necesitaban. "
         "Después organizaron equipos: unos prepararon la tierra, otros sembraron las semillas y otros "
         "registraron el crecimiento de las plantas. Durante las primeras semanas, algunas semillas no "
         "germinaron porque recibieron demasiada agua. El grupo revisó sus registros, modificó el riego y "
         "finalmente consiguió que la mayoría de las plantas creciera. La experiencia les enseñó que observar "
         "los resultados y corregir los errores es parte del aprendizaje.")
COMP = [("¿Qué decidieron hacer los estudiantes?", "crear un huerto escolar", ["huerto"]),
        ("¿Por qué algunas semillas no germinaron?", "recibieron demasiada agua", ["agua"]),
        ("¿Qué hicieron para mejorar los resultados?", "revisaron registros y modificaron el riego", ["riego", "registro"]),
        ("¿Qué enseñanza obtuvieron de la experiencia?", "observar y corregir errores es parte del aprendizaje",
         ["error", "corregir", "observ"])]
MATES = [("Completa la secuencia: 3, 6, 9, 12, ___", ["15"], "15"),
         ("¿Cuál es el resultado de 48 ÷ 6?", ["8"], "8"),
         ("Si un cuaderno cuesta $18, ¿cuánto cuestan 3 cuadernos?", ["54"], "54"),
         ("¿Qué número falta? 5, 10, 20, 40, ___", ["80"], "80"),
         ("Un estudiante leyó 24 páginas el lunes y 17 el martes. ¿Cuántas leyó en total?", ["41"], "41"),
         ("Una actividad empieza a las 10:15 y dura 35 minutos. ¿A qué hora termina?", ["10.50"], "10:50"),
         ("Escribe los números en orden de menor a mayor: 0.8, 0.25, 1.2", ["0.25", "0.8", "1.2"], "0.25, 0.8, 1.2"),
         ("Si 4 paquetes contienen 6 lápices cada uno, ¿cuántos lápices hay en total?", ["24"], "24")]
ORG = ["Olvido qué materiales debo llevar a clase.", "Comienzo una tarea, pero me cuesta terminarla.",
       "Necesito que me recuerden varias veces qué debo hacer.", "Me cuesta organizar una actividad con varios pasos.",
       "Entrego trabajos incompletos por no revisar las instrucciones.",
       "Me cuesta calcular cuánto tiempo necesito para una tarea.",
       "Respondo o actúo antes de pensar qué puede ocurrir.", "Me cuesta volver a una actividad después de una interrupción."]
ORGC = ["Materiales", "Terminar tareas", "Recordatorios", "Organizar pasos", "Instrucciones", "Tiempo", "Impulsividad", "Retomar"]
NIV = {0: "Nunca o casi nunca", 1: "Algunas veces", 2: "Frecuentemente", 3: "Casi siempre"}
APOYOS = ["Instrucciones en pasos cortos", "Repetir o reformular instrucciones", "Apoyo visual / agenda", "Tiempo adicional",
          "Revisión conjunta de tareas", "Lectura guiada o preguntas orales", "Práctica adicional en matemáticas",
          "Menos distractores en el aula", "Entrevista con familia", "Revisión de visión / audición"]
AREAS = {
    "TDAH (presentación inatenta o hiperactiva-impulsiva)": (
        "DSM-5: ≥6 síntomas (≥5 desde los 17 años) de inatención y/o hiperactividad-impulsividad por ≥6 meses, inicio antes de los 12 años, en ≥2 contextos y con afectación funcional. Confirmar con entrevista, escalas con padres y docentes y observación.",
        "sueño, ansiedad o ánimo, visión/audición, estrés familiar, método de estudio, dificultades del aprendizaje",
        "Conners 3 (padres, docentes, autoinforme); prueba de ejecución continua (CPT); d2 (atención selectiva); Stroop; WISC-V: índices de Memoria de Trabajo y Velocidad de Procesamiento"),
    "Dificultad específica del aprendizaje: lectura (dislexia)": (
        "Evaluar precisión, fluidez y comprensión lectora. DSM-5: dificultad ≥6 meses pese a intervención, rendimiento muy por debajo de lo esperado para la edad, sin otra causa que lo explique.",
        "visión/audición, rezago escolar o ausencias, idioma, discapacidad intelectual, falta de práctica",
        "PROLEC-SE-R (lectura en secundaria); WIAT-III o Woodcock-Muñoz (lectura); WISC-V: Comprensión Verbal"),
    "Dificultad específica del aprendizaje: expresión escrita (disgrafía/disortografía)": (
        "Valorar ortografía, gramática, puntuación y claridad/organización del texto, y la grafía en una muestra manuscrita.",
        "motricidad fina, visión, rezago, idioma, falta de práctica",
        "PROESC (procesos de escritura); WIAT-III o Woodcock-Muñoz (expresión escrita); muestra manuscrita"),
    "Dificultad específica del aprendizaje: matemáticas (discalculia)": (
        "Valorar sentido numérico, memoria de operaciones, cálculo y razonamiento; revisar procedimientos y no solo resultados.",
        "ansiedad ante matemáticas, rezago, instrucción previa, errores por inatención",
        "WIAT-III o Woodcock-Muñoz (matemáticas); WISC-V: Razonamiento Fluido y Memoria de Trabajo"),
    "Memoria verbal y procesamiento de instrucciones orales": (
        "Prueba estandarizada de memoria de trabajo/verbal, audiometría y valoración del lenguaje.",
        "audición, ansiedad, idioma, distracciones durante la actividad",
        "WISC-V: Retención de Dígitos y Memoria de Trabajo; audiometría; valoración del lenguaje"),
    "Funciones ejecutivas (planificación, organización, inhibición)": (
        "Escalas de funciones ejecutivas con varios informantes y observación en tareas de varios pasos; hay solapamiento con TDAH.",
        "sueño, estrés, ansiedad, ambiente de estudio",
        "BRIEF-2 (padres, docentes, autoinforme); ENFEN (6 a 12 años); Stroop"),
}

S = st.session_state
ZW = "\u200b"  # carácter invisible para que cada letra tenga una etiqueta única


def mostrar(html, alto):  # st.iframe en versiones nuevas; components.html en las anteriores
    if hasattr(st, "iframe"):
        st.iframe(html, height=alto)
    else:
        components.html(html, height=alto)


def norm(t):
    return "".join(c for c in unicodedata.normalize("NFD", t.lower()) if unicodedata.category(c) != "Mn")


def acierta(texto, esperado):
    toks = re.findall(r"\d+(?:[.:]\d+)?", (texto or "").replace(":", "."))
    return [t[:-3] if t.endswith(".00") else t for t in toks] == esperado


def cuenta(s):
    return (f"<div id=t style='font:600 22px sans-serif;color:#fff;background:#2b6cb0;border-radius:8px;padding:8px;"
            f"text-align:center'></div><script>let s={int(s)};const e=document.getElementById('t');"
            "function f(){e.textContent=s>0?'⏱ '+Math.floor(s/60)+':'+String(s%60).padStart(2,'0'):'¡Tiempo! Presiona terminar';"
            "s--;if(s>=-1)setTimeout(f,1000)}f()</script>")


def pal_js(PALABRAS):
    return (f"<div id=p style='font:700 40px sans-serif;color:#fff;background:#2b6cb0;border-radius:10px;padding:12px;"
          f"text-align:center;min-height:56px'></div><script>const w={PALABRAS};let i=0;const p=document.getElementById('p');"
          "function n(){if(i>=w.length){p.textContent='';return}p.textContent=w[i];"
          "try{const u=new SpeechSynthesisUtterance(w[i]);u.lang='es-MX';speechSynthesis.speak(u)}catch(e){}"
          "i++;setTimeout(n,1500)}n()</script>")


def _r(*xs):
    return [x.split() for x in xs]


_IGNORAR = "sueño, ansiedad o ánimo, visión/audición, estrés, método de estudio, otras dificultades del aprendizaje"
VERS = {
    "ninos": {
        "nombre": "Niñez (menores de 12 años)", "tiempo": 90, "orac": (3, 5),
        "filas": _r("B A D A F E A C B D", "A C E B A D F A C B", "D F A B C A E D A F",
                    "E B D A C F A B D A", "C A F D B A E C A D", "F A B E A C D A B E"),
        "palabras": ["perro", "pelota", "manzana", "silla", "nube"],
        "acciones": ["Nombre arriba", "Círculo abajo a la derecha", "Número 3 dentro del círculo",
                     "Nombre subrayado", "CASA debajo del círculo"],
        "instr2b": "Escribe tu nombre arriba de la hoja. Dibuja un círculo abajo, a la derecha. Escribe el número 3 dentro del círculo. "
                   "Subraya tu nombre. Escribe la palabra CASA debajo del círculo.",
        "texto": ("**EL GATO PELUSA.** Ana tiene un gato llamado Pelusa. Todas las mañanas, Pelusa se sube a la ventana para ver "
                  "a los pájaros. Un día, Ana olvidó cerrar la puerta y Pelusa salió al patio. Ana la buscó por todos lados hasta "
                  "que la encontró durmiendo debajo de un árbol. Desde ese día, Ana siempre revisa la puerta antes de salir."),
        "comp": [("¿Cómo se llama el gato?", "Pelusa", ["pelusa"]),
                 ("¿Qué hacía Pelusa por las mañanas?", "se subía a la ventana a ver los pájaros", ["ventana", "pajaro"]),
                 ("¿Qué olvidó hacer Ana?", "cerrar la puerta", ["puerta", "cerrar"]),
                 ("¿Qué hace Ana desde ese día?", "revisa la puerta antes de salir", ["revis", "antes", "siempre"])],
        "mates": [("Completa la secuencia: 2, 4, 6, 8, ___", ["10"], "10"), ("¿Cuánto es 15 − 7?", ["8"], "8"),
                  ("Si una paleta cuesta $5, ¿cuánto cuestan 3 paletas?", ["15"], "15"),
                  ("¿Qué número falta? 10, 20, 30, ___", ["40"], "40"),
                  ("Tenías 12 canicas y ganaste 9. ¿Cuántas tienes ahora?", ["21"], "21"),
                  ("Son las 3:00 y pasa 1 hora. ¿Qué hora es?", ["4"], "4:00"),
                  ("Ordena de menor a mayor: 9, 3, 6", ["3", "6", "9"], "3, 6, 9"),
                  ("Si hay 3 bolsas con 4 dulces cada una, ¿cuántos dulces hay?", ["12"], "12")],
        "org": ["Se me olvida llevar mis cosas a la escuela.", "Empiezo mi tarea, pero no la termino.",
                "Necesito que me recuerden varias veces lo que debo hacer.", "Me cuesta hacer actividades que tienen varios pasos.",
                "Entrego trabajos incompletos porque no leo bien las instrucciones.",
                "Me cuesta saber cuánto tiempo tardo en hacer una tarea.", "Hago las cosas muy rápido sin pensar.",
                "Me cuesta volver a mi tarea cuando me interrumpen."],
        "p_orac": "Escribe de 3 a 5 oraciones sobre tu mejor día.",
        "p_ayuda": "¿Qué te ayuda a concentrarte o a terminar tus tareas?",
        "p_dif": "¿Hay algo que te cuesta trabajo en la escuela y que quisieras que tu maestro supiera?",
        "tests": [['Conners 3', 'CPT', 'd2', 'Stroop', 'WISC-V'], ['PROLEC-R', 'WIAT-III', 'Woodcock-Muñoz', 'WISC-V'], ['PROESC', 'WIAT-III', 'Woodcock-Muñoz'], ['WIAT-III', 'Woodcock-Muñoz', 'TEDI-MATH', 'WISC-V'], ['WISC-V', 'Audiometría', 'Valoración del lenguaje'], ['BRIEF-2', 'ENFEN', 'Stroop']]},
    "adol": {
        "nombre": "Adolescencia (12 a 15 años)", "tiempo": 120, "orac": (5, 8), "filas": FILAS, "palabras": PALABRAS,
        "acciones": ACCIONES, "texto": TEXTO, "comp": COMP, "mates": MATES, "org": ORG,
        "instr2b": "Escribe tu nombre en la parte superior de la hoja. Dibuja un círculo en la esquina inferior derecha. Escribe el número 8 "
                   "dentro del círculo. Subraya tu nombre. Escribe la palabra ESCUELA debajo del círculo.",
        "p_orac": "Escribe un texto de 5 a 8 oraciones sobre una situación en la que tuviste que resolver un problema.",
        "p_ayuda": "¿Qué te ayuda a concentrarte o terminar tus tareas?",
        "p_dif": "¿Hay algo que te dificulte aprender o participar en clase y que quisieras que la escuela supiera?",
        "tests": [['Conners 3', 'CPT', 'd2', 'Stroop', 'WISC-V'], ['PROLEC-SE-R', 'WIAT-III', 'Woodcock-Muñoz', 'WISC-V'], ['PROESC', 'WIAT-III', 'Woodcock-Muñoz'], ['WIAT-III', 'Woodcock-Muñoz', 'WISC-V'], ['WISC-V', 'Audiometría', 'Valoración del lenguaje'], ['BRIEF-2', 'Stroop']]},
    "adulto": {
        "nombre": "Adultos (16 años o más)", "tiempo": 120, "orac": (8, 12),
        "filas": _r("V A H R K A N X V A R H K A", "K H A V R N A H X A K V R N", "N R V A H K X A R V A H N K",
                    "H A K N V R A X H K A V N R", "R V A X H A N K V R A H X A", "X K N A R V H A K X N A V H",
                    "A H R V X K A N H R A V K X", "V N K H A R X V A H K A R N"),
        "palabras": ["reloj", "puente", "naranja", "cuaderno", "cortina"],
        "acciones": ["Nombre arriba", "Cuadrado abajo a la izquierda", "Número 7 dentro del cuadrado",
                     "Nombre subrayado", "TRABAJO encima del cuadrado"],
        "instr2b": "Escribe tu nombre en la parte superior de la hoja. Dibuja un cuadrado en la esquina inferior izquierda. Escribe el número 7 "
                   "dentro del cuadrado. Subraya tu nombre. Escribe la palabra TRABAJO encima del cuadrado.",
        "texto": ("**EL CAMBIO DE HORARIO.** Una empresa mediana decidió cambiar el horario de trabajo para reducir el tiempo que sus "
                  "empleados pasaban en el tráfico. Antes de aplicar la medida, la gerencia consultó a los equipos y comparó los costos. "
                  "Durante el primer mes, la productividad bajó porque varios trabajadores no lograron adaptar sus rutinas familiares. "
                  "La empresa revisó las quejas, ajustó los horarios de dos áreas y ofreció trabajar un día desde casa. Tras tres meses, "
                  "la rotación de personal disminuyó y los empleados reportaron menos cansancio. La gerencia concluyó que los cambios "
                  "funcionan mejor cuando se revisan con los afectados y se corrigen a tiempo."),
        "comp": [("¿Para qué decidió la empresa cambiar el horario?", "reducir el tiempo en el tráfico", ["trafico", "traslado", "transporte"]),
                 ("¿Por qué bajó la productividad el primer mes?", "varios trabajadores no lograron adaptar sus rutinas familiares", ["rutina", "adapt", "familia"]),
                 ("¿Qué medidas tomó la empresa tras revisar las quejas?", "ajustó horarios de dos áreas y permitió un día desde casa", ["ajust", "casa", "home"]),
                 ("¿Qué concluyó la gerencia?", "los cambios funcionan mejor si se revisan con los afectados y se corrigen a tiempo", ["revis", "corrig", "afectad", "consult"])],
        "mates": [("Completa la secuencia: 4, 9, 16, 25, ___", ["36"], "36"), ("Calcula el 15% de 240", ["36"], "36"),
                  ("Un artículo cuesta $850 y tiene 20% de descuento. ¿Cuánto pagas?", ["680"], "680"),
                  ("¿Qué número sigue? 2, 6, 18, 54, ___", ["162"], "162"),
                  ("Compraste tres productos de $128, $245 y $87. ¿Cuánto gastaste en total?", ["460"], "460"),
                  ("Una reunión empieza a las 9:40 y dura 1 hora con 35 minutos. ¿A qué hora termina?", ["11.15"], "11:15"),
                  ("Ordena de menor a mayor: 0.75, 0.5, 0.125", ["0.125", "0.5", "0.75"], "0.125, 0.5, 0.75"),
                  ("Si 7 personas pagan $1,260 en total por partes iguales, ¿cuánto paga cada una?", ["180"], "180")],
        "org": ["Olvido qué materiales o documentos necesito llevar.", "Comienzo tareas, pero me cuesta terminarlas.",
                "Necesito recordatorios frecuentes para cumplir mis pendientes.", "Me cuesta organizar actividades de varios pasos.",
                "Entrego trabajos incompletos por no revisar bien las instrucciones.",
                "Me cuesta calcular cuánto tiempo necesito para una tarea.",
                "Actúo o respondo antes de pensar en las consecuencias.", "Me cuesta retomar una actividad después de una interrupción."],
        "p_orac": "Escribe un texto de 8 a 12 oraciones sobre una situación reciente de estudio, trabajo o vida diaria en la que tuviste que resolver un problema.",
        "p_ayuda": "¿Qué te ayuda a concentrarte o a terminar tus pendientes?",
        "p_dif": "¿Hay algo que dificulte tu estudio o trabajo y que quisieras mencionar?",
        "tests": [['ASRS', 'CAARS', 'Conners 3', 'CPT', 'd2', 'Stroop', 'WAIS-IV'], ['WIAT-III', 'Woodcock-Muñoz', 'WAIS-IV'], ['WIAT-III', 'Woodcock-Muñoz'], ['WIAT-III', 'Woodcock-Muñoz', 'WAIS-IV'], ['WAIS-IV', 'Audiometría', 'Valoración del lenguaje'], ['BRIEF-A', 'BRIEF-2', 'Stroop', 'Torre de Londres']]},
}

FICHAS = {
    "Conners 3": ("Síntomas de TDAH y problemas asociados (conducta, aprendizaje, ansiedad)", "6 a 18 años", "Padres, docentes y autoinforme (desde 8 años)", "Licencia comercial; requiere formación"),
    "CPT": ("Atención sostenida, impulsividad y tiempo de reacción en computadora", "Desde 6 años, según la versión", "Lo aplica el profesional", "Licencia comercial (hay varias versiones)"),
    "d2": ("Atención selectiva y concentración sostenida (tachado de letras)", "Desde 8 años", "Lo aplica el profesional (unos 8 min, en papel)", "Licencia comercial"),
    "Stroop": ("Control inhibitorio, atención y flexibilidad", "Versiones infantil y adulta", "Lo aplica el profesional (unos 5 min)", "Licencia comercial"),
    "WISC-V": ("Perfil cognitivo: Comprensión Verbal, Visoespacial, Razonamiento Fluido, Memoria de Trabajo y Velocidad de Procesamiento", "6 a 16 años 11 meses", "Lo aplica un psicólogo (60 a 90 min)", "Licencia comercial; uso restringido a psicólogos"),
    "WAIS-IV": ("Perfil cognitivo: Comprensión Verbal, Razonamiento Perceptual, Memoria de Trabajo y Velocidad de Procesamiento", "16 a 90 años", "Lo aplica un psicólogo (60 a 90 min)", "Licencia comercial; uso restringido a psicólogos"),
    "PROLEC-R": ("Procesos lectores: decodificación y comprensión", "6 a 12 años", "Lo aplica el profesional", "Licencia comercial"),
    "PROLEC-SE-R": ("Procesos lectores en secundaria y bachillerato", "12 a 16 años", "Lo aplica el profesional", "Licencia comercial"),
    "PROESC": ("Procesos de escritura: ortografía, reglas y composición", "8 a 15 años", "Lo aplica el profesional", "Licencia comercial"),
    "WIAT-III": ("Aprovechamiento: lectura, escritura, matemáticas y lenguaje oral", "4 a 50 años", "Lo aplica el profesional", "Licencia comercial"),
    "Woodcock-Muñoz": ("Batería de aprovechamiento y habilidades cognitivas en español", "Desde 2 años hasta adultos", "Lo aplica el profesional", "Licencia comercial"),
    "TEDI-MATH": ("Competencias numéricas básicas", "4 a 8 años", "Lo aplica el profesional", "Licencia comercial"),
    "BRIEF-2": ("Funciones ejecutivas en la vida diaria", "5 a 18 años", "Padres, docentes y autoinforme (desde 11 años)", "Licencia comercial"),
    "BRIEF-A": ("Funciones ejecutivas en adultos", "18 a 90 años", "Autoinforme e informante cercano", "Licencia comercial"),
    "ENFEN": ("Funciones ejecutivas: fluidez, atención, planificación e inhibición", "6 a 12 años", "Lo aplica el profesional", "Licencia comercial"),
    "ASRS": ("Cribado de síntomas de TDAH en adultos (OMS)", "18 años o más", "Autoinforme", "Cribado: no diagnostica; consulta los términos de uso"),
    "CAARS": ("Síntomas de TDAH en adultos", "18 años o más", "Autoinforme y observador", "Licencia comercial"),
    "Torre de Londres": ("Planificación y resolución de problemas", "Versiones infantil y adulta", "Lo aplica el profesional", "Licencia comercial"),
    "Audiometría": ("Descartar pérdida auditiva", "Todas las edades", "Audiólogo o profesional de salud auditiva", "Servicio de salud"),
    "Valoración del lenguaje": ("Lenguaje comprensivo y expresivo", "Según la edad", "Psicólogo o terapeuta del lenguaje", "Según el instrumento elegido"),
    "SDQ": ("Fortalezas y dificultades emocionales y conductuales", "4 a 17 años", "Padres, docentes y autoinforme (desde 11 años)", "Uso gratuito con condiciones del autor; consulta sus términos"),
    "CBCL": ("Problemas emocionales y conductuales", "6 a 18 años", "Padres", "Licencia comercial (ASEBA)"),
}
OPC = {0: "Para nada", 1: "Varios días", 2: "Más de la mitad de los días", 3: "Casi todos los días"}
PHQ9 = ["Poco interés o placer en hacer las cosas", "Se ha sentido decaído(a), deprimido(a) o sin esperanzas",
        "Dificultad para quedarse o permanecer dormido(a), o ha dormido demasiado", "Se ha sentido cansado(a) o con poca energía",
        "Con poco apetito o ha comido en exceso",
        "Se ha sentido mal consigo mismo(a), o que es un fracaso, o que ha quedado mal con usted mismo(a) o con su familia",
        "Dificultad para concentrarse en cosas como leer o ver televisión",
        "Se ha movido o hablado tan lentamente que otras personas podrían notarlo, o al contrario, ha estado tan inquieto(a) que se ha movido mucho más de lo normal",
        "Pensamientos de que estaría mejor muerto(a) o de lastimarse de alguna manera"]
GAD7 = ["Se ha sentido nervioso(a), ansioso(a) o con los nervios de punta", "No ha podido dejar de preocuparse o controlar la preocupación",
        "Se ha preocupado demasiado por motivos diferentes", "Ha tenido dificultad para relajarse",
        "Se ha sentido tan inquieto(a) que no ha podido quedarse quieto(a)", "Se ha molestado o irritado fácilmente",
        "Ha sentido miedo como si algo terrible fuera a pasar"]
R_PHQ = [(4, "mínimo"), (9, "leve"), (14, "moderado"), (19, "moderadamente grave"), (27, "grave")]
R_GAD = [(4, "mínimo"), (9, "leve"), (14, "moderado"), (21, "grave")]


def ficha(t):
    mide, edad_, quien, disp = FICHAS[t]
    with st.expander(f"🔒 {t}"):
        st.markdown(f"**Mide:** {mide}  \n**Edades:** {edad_}  \n**Quién responde:** {quien}  \n**Disponibilidad:** {disp}")
        st.link_button("Buscar información y dónde obtenerla", "https://www.google.com/search?q=" + quote(f"{t} prueba psicológica"))


def aplicar(nombre, items, pref, rangos):
    with st.expander(f"🟢 {nombre} · aplicar aquí"):
        st.caption("Durante las últimas 2 semanas, ¿con qué frecuencia le han molestado los siguientes problemas?")
        vals = [st.radio(t, list(OPC), index=None, format_func=lambda v: OPC[v], horizontal=True, key=f"{pref}{i}")
                for i, t in enumerate(items)]
        n = sum(v is not None for v in vals)
        if n < len(items):
            st.info(f"Respondidas {n} de {len(items)}")
            return None
        total = sum(vals)
        nivel = next(nv for lim, nv in rangos if total <= lim)
        st.markdown(f"**Puntaje: {total} · {nivel}**")
        return total, nivel, vals


def _tabla(doc, enc, filas):
    t = doc.add_table(rows=1, cols=len(enc))
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, enc):
        c.text = h
    for f in filas:
        for c, v in zip(t.add_row().cells, f):
            c.text = str(v)
    doc.add_paragraph()


def _etiqueta(doc, etiqueta, texto):
    p = doc.add_paragraph()
    p.add_run(etiqueta + " ").bold = True
    p.add_run(texto)


def crear_docx(d):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Reporte de exploración educativa", 0)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in d["datos"]:
        c = t.add_row().cells
        c[0].text, c[1].text = k, str(v)
    doc.add_heading("Resumen", 1)
    for x in d["resumen"]:
        doc.add_paragraph(x)
    doc.add_heading("Resultados por actividad", 1)
    _tabla(doc, ["Actividad", "Resultado"], d["resultados"])
    doc.add_heading("Perfil por dominios", 1)
    doc.add_paragraph("Referencia propia y no normativa (menos de 40 bajo, 40 a 69 medio, 70 o más adecuado).")
    _tabla(doc, ["Dominio", "Puntaje", "Lectura"], d["dominios"])
    doc.add_heading("Actividad 4. Texto del estudiante", 1)
    doc.add_paragraph(d["txt4"] or "-")
    doc.add_heading("Actividad 6. Organización de tareas", 1)
    _tabla(doc, ["Conducta", "Frecuencia"], d["org"])
    _etiqueta(doc, "Le ayuda:", d["ayuda"] or "-")
    _etiqueta(doc, "Le dificulta:", d["dif"] or "-")
    doc.add_heading("Áreas a explorar (hipótesis de trabajo)", 1)
    if not d["areas"]:
        doc.add_paragraph("Sin áreas que cumplan los criterios. Esto no descarta nada.")
    for a in d["areas"]:
        doc.add_heading(f"{a['n']} · indicio {a['niv']}", 2)
        for x in a["indicios"]:
            doc.add_paragraph(x, style="List Bullet")
        _etiqueta(doc, "Criterios:", a["criterios"])
        _etiqueta(doc, "Descartar primero:", a["descartar"])
        _etiqueta(doc, "Pruebas sugeridas:", a["pruebas"])
    if d["emo"]:
        doc.add_heading("Tamizaje emocional", 1)
        for x in d["emo"]:
            doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("Observaciones del aplicador", 1)
    doc.add_paragraph(d["obs"] or "-")
    doc.add_heading("Apoyos que se probarán", 1)
    for x in d["apoyos"] or ["-"]:
        doc.add_paragraph(x, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


tab_ev, tab_re = st.tabs(["📝 Evaluación", "📊 Resultados"])

# ======================= EVALUACIÓN =======================
with tab_ev:
    st.title("Cuadernillo de exploración educativa")
    st.caption("No es un examen: si no sabes una respuesta, déjala en blanco y continúa.")
    c1, c2 = st.columns(2)
    c1.text_input("Nombre", key="w_nombre")
    edad_in = c2.number_input("Edad", 3, 99, 12, key="w_edad")
    c1.date_input("Fecha de aplicación", key="w_fecha")
    c2.text_input("Aplicador(a)", key="w_aplicador")
    clave = "ninos" if edad_in < 12 else "adol" if edad_in < 16 else "adulto"
    if S.get("_ver") != clave:  # al cambiar de versión se borran las respuestas anteriores
        pref = ("w_r", "w_c", "w_m", "w_o", "w_2a", "w_a4", "w_ayuda", "w_dif", "v_")
        for k in [k for k in S if k.startswith(pref) or k in ("_t1", "_f1", "_oyo")]:
            del S[k]
        S["_ver"] = clave
    V = VERS[clave]
    FILAS, PALABRAS, ACCIONES, TEXTO, COMP, MATES, ORG = (V[k] for k in ("filas", "palabras", "acciones", "texto", "comp", "mates", "org"))
    TIEMPO1, (ORMIN, ORMAX) = V["tiempo"], V["orac"]
    NF, NA = len(FILAS), sum(r.count("A") for r in FILAS)
    st.caption(f"Versión: {V['nombre']}")

    st.header("Actividad 1. Atención visual")
    st.write(f"Durante {TIEMPO1 / 60:g} minutos, **toca todas las letras A** que encuentres para tacharlas. No toques ninguna otra letra. "
             "Trabaja de izquierda a derecha y de arriba hacia abajo.")
    if st.button(f"▶ Comenzar ({TIEMPO1 / 60:g} minutos)", disabled="_t1" in S):
        S["_t1"] = time.time()
        st.rerun()
    if "_t1" in S:
        ahora = time.time()
        if "_f1" not in S and ahora - S["_t1"] >= TIEMPO1:
            S["_f1"] = S["_t1"] + TIEMPO1
        bloq = "_f1" in S
        if not bloq:
            mostrar(cuenta(TIEMPO1 - (ahora - S["_t1"])), 50)
        for r, fila in enumerate(FILAS):
            st.pills(f"Fila {r + 1}", [f"{c}{ZW * i}" for i, c in enumerate(fila)], selection_mode="multi",
                     key=f"w_r{r}", disabled=bloq, label_visibility="collapsed")
        if bloq:
            st.success("Actividad 1 terminada.")
        elif st.button("Terminé la actividad 1"):
            S["_f1"] = time.time()
            st.rerun()

    st.header("Actividad 2. Escuchar y recordar")
    st.write("Presiona el botón **una sola vez** y presta atención a las palabras. Después escribe todas las que recuerdes, en cualquier orden.")
    if st.button("🔊 Escuchar las palabras", disabled=S.get("_oyo", False)):
        S["_oyo"] = True
        mostrar(pal_js(PALABRAS), 90)
        time.sleep(len(PALABRAS) * 1.5 + 2.5)  # el campo sigue bloqueado hasta que termina el audio
        st.rerun()
    st.text_area("Palabras que recuerdas (se habilita al terminar el audio)", key="w_2a", disabled=not S.get("_oyo"))

    st.header("Actividad 3. Comprensión lectora")
    st.info(TEXTO)
    for i, (q, _, _) in enumerate(COMP):
        st.text_area(f"{i + 1}. {q}", key=f"w_c{i}")

    st.header("Actividad 4. Expresión escrita")
    st.text_area(V["p_orac"] + " "
                 "Explica qué ocurrió, qué hiciste y cómo terminó.", height=220, key="w_a4")

    st.header("Actividad 5. Razonamiento y matemáticas")
    st.caption("Sin calculadora.")
    for i, (q, _, _) in enumerate(MATES):
        st.text_input(f"{i + 1}. {q}", key=f"w_m{i}")

    st.header("Actividad 6. Cómo organizo mis tareas")
    st.caption("Piensa en lo que te sucede normalmente durante las últimas cuatro semanas.")
    for i, t in enumerate(ORG):
        st.radio(t, list(NIV), index=None, format_func=lambda v: f"{v} · {NIV[v]}", horizontal=True, key=f"w_o{i}")
    st.text_area(V["p_ayuda"], key="w_ayuda")
    st.text_area(V["p_dif"], key="w_dif")
    if st.button("✅ Terminé la evaluación", type="primary"):
        st.success("¡Gracias por participar! Avisa a tu aplicador(a).")

# ======================= RESULTADOS =======================
with tab_re:
    if not S.get("_ok"):
        p = st.text_input("PIN del profesional", type="password", key="_pin_in")
        if p == PIN:
            S["_ok"] = True
            st.rerun()
        if p:
            st.error("PIN incorrecto.")
    else:
        if st.button("🔒 Bloquear resultados"):
            S["_ok"] = False
            st.rerun()
        nombre, edad, fecha, aplic = S.get("w_nombre", ""), S.get("w_edad", ""), S.get("w_fecha", ""), S.get("w_aplicador", "")
        st.subheader(f"{nombre or 'Sin nombre'} · {edad} años · {V['nombre']}")
        st.caption(f"Fecha: {fecha} · Aplicador(a): {aplic or '-'}")
        with st.expander("⚙️ Umbrales de las áreas a explorar"):
            for k, lab in ULAB.items():
                U[k] = st.number_input(lab, 0, 25, U[k], key=f"u_{k}")

        # Act 1
        sel = [[x.count(ZW) for x in (S.get(f"w_r{r}") or [])] for r in range(NF)]
        hits = [sum(FILAS[r][i] == "A" for i in s) for r, s in enumerate(sel)]
        fps = [len(s) - h for s, h in zip(sel, hits)]
        aciertos, falsos = sum(hits), sum(fps)
        omisiones = NA - aciertos
        t1 = f"{S['_f1'] - S['_t1']:.0f} s" if "_f1" in S else "no realizada"
        # Act 2A
        toks = list(dict.fromkeys(re.findall(r"[a-z]+", norm(S.get("w_2a", "")))))
        objetivo = {norm(p): p for p in PALABRAS}
        recordadas = [objetivo[t] for t in toks if t in objetivo]
        extras = [t for t in toks if t not in objetivo]
        # Act 5
        mates_ok = [acierta(S.get(f"w_m{i}", ""), MATES[i][1]) for i in range(8)]
        # Act 4
        txt4 = S.get("w_a4", "")
        n_or = len([s for s in re.split(r"[.!?]+", txt4) if s.strip()])
        # Act 6
        org = {t: S.get(f"w_o{i}") for i, t in enumerate(ORG)}
        org_alto = [t for t, v in org.items() if v is not None and v >= 2]

        st.markdown("#### Verificación del profesional")
        st.markdown("**Actividad 2B (en papel):** marca las acciones ejecutadas correctamente")
        st.caption("Instrucciones de la 2B (leer una sola vez): " + V["instr2b"])
        b_ap = st.checkbox("Se aplicó la actividad 2B", key="v_bap")
        acciones_ok = [a for i, a in enumerate(ACCIONES) if st.checkbox(a, key=f"v_b{i}")]
        st.markdown("**Actividad 3:** sugerencia automática por palabras clave; corrige con tu criterio")
        compr_ok = []
        for i, (q, clave, kw) in enumerate(COMP):
            ans = S.get(f"w_c{i}", "")
            auto = any(k in norm(ans) for k in kw)
            if st.checkbox(f"{q} → «{ans or 'sin respuesta'}» (clave: {clave})", value=auto,
                           key=f"v_c{i}_{zlib.crc32(ans.encode())}"):
                compr_ok.append(q)
        escr_flag = st.checkbox("Actividad 4: texto poco comprensible o muy desorganizado", key="v_escr")

        resumen = pd.DataFrame([
            ["1. Atención visual", f"{aciertos}/{NA} aciertos · {omisiones} omisiones · {falsos} falsos positivos · {t1}"],
            ["2A. Recuerdo inmediato", f"{len(recordadas)}/5" + (f" · otras palabras: {', '.join(extras)}" if extras else "")],
            ["2B. Instrucciones", f"{len(acciones_ok)}/5"],
            ["3. Comprensión lectora", f"{len(compr_ok)}/4"],
            ["4. Expresión escrita", f"{n_or} oraciones · {len(txt4.split())} palabras"],
            ["5. Matemáticas", f"{sum(mates_ok)}/8"],
            ["6. Conductas frecuentes o casi siempre", f"{len(org_alto)}/8"]], columns=["Actividad", "Resultado"])
        st.markdown("#### Resumen")
        st.table(resumen)
        st.markdown("**Porcentaje logrado por actividad**")
        st.bar_chart(pd.DataFrame({"%": [aciertos / NA * 100, len(recordadas) / 5 * 100, len(acciones_ok) / 5 * 100,
                                          len(compr_ok) / 4 * 100, sum(mates_ok) / 8 * 100]},
                                  index=["1 Atención", "2A Recuerdo", "2B Instrucciones", "3 Comprensión", "5 Matemáticas"]))
        st.markdown(f"**Actividad 1: aciertos y falsos positivos por fila** (en total hay {NA} letras A)")
        st.bar_chart(pd.DataFrame({"Aciertos": hits, "Falsos positivos": fps}, index=[f"Fila {r + 1}" for r in range(NF)]))
        st.markdown("**Actividad 6: frecuencia por conducta (0 a 3)**")
        st.bar_chart(pd.DataFrame({"Nivel": [v if v is not None else 0 for v in org.values()]}, index=ORGC))
        st.markdown("**Actividad 5: respuestas**")
        st.dataframe(pd.DataFrame({"Pregunta": [m[0] for m in MATES], "Respuesta": [S.get(f"w_m{i}", "") for i in range(8)],
                                   "Clave": [m[2] for m in MATES], "✓": ["✔" if o else "✘" for o in mates_ok]}),
                     hide_index=True)
        st.markdown(f"**Actividad 4: texto del estudiante** (se esperan {ORMIN} a {ORMAX} oraciones)")
        st.write(txt4 or "_Sin respuesta_")
        st.markdown(f"**Actividad 6, respuestas abiertas.** Le ayuda: {S.get('w_ayuda') or '-'} · Le dificulta: {S.get('w_dif') or '-'}")

        # ---- Perfil por dominios (referencia propia, no normativa) ----
        t1s = S["_f1"] - S["_t1"] if "_f1" in S else None
        mid = NF // 2
        h1, h2 = sum(hits[:mid]), sum(hits[mid:])
        suma = lambda ids: sum(mates_ok[i] for i in ids)
        calc, razon = suma([1, 2, 4, 7]), suma([0, 3, 5, 6])
        vals = [v for v in org.values() if v is not None]
        hizo3 = any(S.get(f"w_c{i}", "").strip() for i in range(4))
        hizo5 = any(S.get(f"w_m{i}", "").strip() for i in range(8))
        dom = {"Atención sostenida (eficiencia)": max(0, (aciertos - falsos) / NA * 100) if "_t1" in S else None,
               "Memoria verbal inmediata": len(recordadas) / 5 * 100 if S.get("_oyo") else None,
               "Seguimiento de instrucciones": len(acciones_ok) / 5 * 100 if b_ap else None,
               "Comprensión lectora": len(compr_ok) / 4 * 100 if hizo3 else None,
               "Cálculo básico": calc / 4 * 100 if hizo5 else None,
               "Razonamiento matemático": razon / 4 * 100 if hizo5 else None,
               "Autorregulación (autoinforme)": 100 - sum(vals) / (3 * len(vals)) * 100 if vals else None}
        dd = pd.DataFrame({"Dominio": list(dom), "Puntaje %": [None if v is None else round(v) for v in dom.values()]})
        dd["Lectura descriptiva"] = dd["Puntaje %"].map(
            lambda v: "No evaluado" if pd.isna(v) else "Bajo" if v < 40 else "Medio" if v < 70 else "Adecuado")
        st.markdown("#### Perfil por dominios")
        st.caption("Porcentaje de logro con referencia propia y no normativa (<40 bajo, 40-69 medio, ≥70 adecuado). "
                   "Para valores estandarizados usa las pruebas sugeridas más abajo.")
        st.dataframe(dd, hide_index=True)
        if dd["Puntaje %"].notna().any():
            st.bar_chart(dd.dropna().set_index("Dominio")["Puntaje %"])
        if "_t1" in S:
            m = st.columns(4)
            m[0].metric("Omisiones", f"{omisiones / NA:.0%}")
            m[1].metric("Falsos positivos", falsos)
            m[2].metric("Aciertos 1ª vs 2ª mitad", f"{h1} vs {h2}")
            m[3].metric("Tiempo", t1)

        # ---- Pistas en respuestas abiertas ----
        abiertas = norm(f"{S.get('w_ayuda', '')} {S.get('w_dif', '')}")
        CLAVES = {"Emocional / social": ["miedo", "nervio", "ansied", "trist", "llor", "enoj", "burla", "bullying", "acoso", "estres", "preocup"],
                  "Sueño / cansancio": ["sueno", "cansad", "dormir"],
                  "Atención / ambiente": ["distra", "ruido", "platic", "celular", "aburr"],
                  "Comprensión / ritmo": ["no entiendo", "no comprendo", "muy rapido", "me cuesta leer", "no puedo leer"]}
        pistas = {k: [w for w in v if w in abiertas] for k, v in CLAVES.items()}
        pistas = {k: v for k, v in pistas.items() if v}

        # ---- Indicios por área ----
        K = list(AREAS)
        ind = {n: [] for n in K}
        if "_t1" in S:
            if omisiones >= U["omisiones"] * NA / 25: ind[K[0]].append(f"Omisiones: {omisiones}/{NA}")
            if falsos >= U["falsos_pos"]: ind[K[0]].append(f"Falsos positivos: {falsos} (posible impulsividad)")
            if h1 - h2 >= U["fatiga"]: ind[K[0]].append(f"El desempeño cae en la segunda mitad ({h1} vs {h2} aciertos)")
            if t1s is not None and t1s < 60 and omisiones >= U["omisiones"] * NA / 25:
                ind[K[0]].append(f"Terminó en {t1s:.0f} s con muchas omisiones (posible impulsividad)")
        if S.get("_oyo") and len(recordadas) <= U["recuerdo"]:
            ind[K[0]].append(f"Recuerdo inmediato bajo: {len(recordadas)}/5")
            ind[K[4]].append(f"Recuerdo de palabras: {len(recordadas)}/5")
        intr = [t for t in extras if len(t) >= 4]
        if len(intr) >= 2: ind[K[4]].append(f"Intrusiones (palabras no presentadas): {', '.join(intr)}")
        if b_ap and len(acciones_ok) <= U["instrucciones"]: ind[K[4]].append(f"Instrucciones de 5 pasos: {len(acciones_ok)}/5")
        if len(org_alto) >= U["org_items"]: ind[K[0]].append(f"Autoinforme: {len(org_alto)} de 8 conductas de organización frecuentes")
        GRUPOS = {"Planificación y organización": [0, 3, 5], "Persistencia e iniciativa": [1, 2],
                  "Inhibición y monitoreo": [4, 6], "Flexibilidad al retomar tareas": [7]}
        for g, ids in GRUPOS.items():
            vs = [org[ORG[i]] for i in ids if org[ORG[i]] is not None]
            if vs and sum(vs) / len(vs) >= 2: ind[K[5]].append(f"{g}: promedio {sum(vs) / len(vs):.1f} de 3")
        if hizo3 and len(compr_ok) <= U["comprension"]: ind[K[1]].append(f"Comprensión lectora: {len(compr_ok)}/4")
        if sum(not S.get(f"w_c{i}", "").strip() for i in range(4)) >= 2 and hizo3:
            ind[K[1]].append("Dejó sin respuesta 2 o más preguntas de comprensión")
        if "Comprensión / ritmo" in pistas: ind[K[1]].append("Menciona dificultad para entender o leer en sus respuestas abiertas")
        pal4 = re.findall(r"[a-zñ]+", norm(txt4))
        ttr = len(set(pal4)) / len(pal4) if pal4 else 0
        if txt4.strip():
            if n_or < ORMIN: ind[K[2]].append(f"Texto breve: {n_or} oraciones (se pedían {ORMIN} a {ORMAX})")
            if len(pal4) >= 30 and ttr < 0.5: ind[K[2]].append(f"Vocabulario repetitivo (diversidad léxica {ttr:.2f})")
            if len(pal4) / max(n_or, 1) > 30: ind[K[2]].append(f"Oraciones muy largas o sin puntuar ({len(pal4) / max(n_or, 1):.0f} palabras por oración)")
        if escr_flag: ind[K[2]].append("Texto poco comprensible o desorganizado (según el profesional)")
        if hizo5:
            if sum(mates_ok) <= U["mates"]: ind[K[3]].append(f"Matemáticas: {sum(mates_ok)}/8 correctas")
            if calc <= 1: ind[K[3]].append(f"Cálculo básico (÷, ×, +): {calc}/4")
            if razon <= 1: ind[K[3]].append(f"Razonamiento y patrones: {razon}/4")
        niv = lambda n: ["Sin indicios", "Débil", "Moderado", "Fuerte"][min(len(ind[n]), 3)]
        areas = [(n, ind[n]) for n in K if ind[n]]

        st.markdown("#### Indicios por área")
        st.caption("Hipótesis de trabajo para tu criterio profesional, no un diagnóstico. El nivel solo refleja cuántos indicios coinciden "
                   "(1 débil, 2 moderado, 3 o más fuerte). Este cuadernillo no evalúa ansiedad, depresión, TEA, lenguaje ni capacidad intelectual.")
        st.table(pd.DataFrame({"Área": K, "Indicio": [niv(n) for n in K], "Coincidencias": [len(ind[n]) for n in K]}))
        for n, indicios in areas:
            with st.container(border=True):
                st.markdown(f"##### {n} · {niv(n)}")
                st.markdown("**Por qué aparece:**\n" + "\n".join(f"- {x}" for x in indicios))
                st.markdown(f"**Criterios y cómo confirmar:** {AREAS[n][0]}")
                st.markdown(f"**Descartar primero:** {AREAS[n][1]}")
                st.markdown("**Pruebas recomendadas** (abre una para ver su ficha)")
                for t in V["tests"][K.index(n)]:
                    ficha(t)
        if pistas:
            st.markdown("**Pistas en sus respuestas abiertas** (para explorar en entrevista)")
            for k, v in pistas.items():
                st.write(f"- {k}: " + ", ".join(v))
            if "Emocional / social" in pistas:
                st.caption("Para lo emocional y conductual se sugiere SDQ o CBCL con padres y docentes.")
        st.caption("Verifica edad, baremos para México y licencias de cada prueba antes de usarla.")

        st.markdown("#### Tamizaje emocional (gratuito, se aplica aquí)")
        res_emo = []
        if edad >= 12:
            st.caption("El cuadernillo no evalúa ánimo ni ansiedad. PHQ-9 y GAD-7 son escalas de dominio público para adultos; "
                       "de 12 a 17 años úsalas solo como exploración y confirma con instrumentos para adolescentes.")
            r = aplicar("PHQ-9 (depresión)", PHQ9, "v_p9_", R_PHQ)
            if r:
                res_emo.append(f"- PHQ-9: {r[0]} ({r[1]})")
                if r[2][8] > 0:
                    st.error("El ítem 9 tiene respuesta mayor que 0: valora el riesgo de forma directa y sigue tu protocolo de intervención en crisis.")
                    res_emo.append("- PHQ-9 ítem 9 mayor que 0: valorar riesgo de inmediato")
            r = aplicar("GAD-7 (ansiedad)", GAD7, "v_g7_", R_GAD)
            if r:
                res_emo.append(f"- GAD-7: {r[0]} ({r[1]})")
        else:
            st.caption("Para menores de 12 años se sugieren estos cuestionarios con padres y docentes; no se aplican aquí.")
            ficha("SDQ")
            ficha("CBCL")

        st.markdown("#### Observaciones y apoyos")
        obs = st.text_area("Observaciones del aplicador (atención, conducta, fortalezas, factores que influyeron)", key="v_obs")
        apoyos = st.multiselect("Apoyos inmediatos que se probarán", APOYOS, key="v_apoyos")

        bajos = [r["Dominio"] for _, r in dd.iterrows() if r["Lectura descriptiva"] == "Bajo"]
        res = [f"{nombre or 'El evaluado'} ({edad} años) respondió la versión «{V['nombre']}» el {fecha}."]
        res.append("Áreas con indicios para explorar: " + ", ".join(f"{n} (indicio {niv(n).lower()})" for n, _ in areas) + "."
                   if areas else "Ninguna área cumplió los criterios de indicio; esto no descarta nada.")
        if bajos:
            res.append("Dominios con puntaje bajo (referencia no normativa): " + ", ".join(bajos) + ".")
        if res_emo:
            res.append("Tamizaje emocional: " + "; ".join(x.lstrip("- ") for x in res_emo) + ".")
        res.append("Este resumen reúne hipótesis de trabajo para el criterio del profesional; no es un diagnóstico.")
        d = {"datos": [("Nombre", nombre or "-"), ("Edad", f"{edad} años"), ("Versión aplicada", V["nombre"]),
                       ("Fecha de aplicación", str(fecha)), ("Aplicador(a)", aplic or "-")],
             "resumen": res,
             "resultados": [(r.Actividad, r.Resultado) for r in resumen.itertuples()],
             "dominios": [(r["Dominio"], "sin evaluar" if pd.isna(r["Puntaje %"]) else f"{int(r['Puntaje %'])} %", r["Lectura descriptiva"])
                          for _, r in dd.iterrows()],
             "txt4": txt4, "ayuda": S.get("w_ayuda"), "dif": S.get("w_dif"),
             "org": [(t, "sin responder" if v is None else NIV[v]) for t, v in org.items()],
             "areas": [{"n": n, "niv": niv(n).lower(), "indicios": ind_, "criterios": AREAS[n][0], "descartar": AREAS[n][1],
                        "pruebas": ", ".join(V["tests"][K.index(n)])} for n, ind_ in areas],
             "emo": [x.lstrip("- ") for x in res_emo], "obs": obs, "apoyos": apoyos}
        if Document is None:
            st.warning("Para descargar el reporte en Word agrega python-docx al archivo requirements.txt.")
        else:
            st.download_button("⬇ Descargar reporte (Word)", crear_docx(d), f"reporte_{nombre or 'estudiante'}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
